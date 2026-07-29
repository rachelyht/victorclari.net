"""Geometry and packaging checks on the composed deck — the preview cannot catch these."""

from __future__ import annotations

import zipfile

from app.deck import SLIDE_H_IN, SLIDE_W_IN, Clip, compose_deck, image_key
from conftest import needs_ffmpeg, plan_for
from pptx import Presentation
from pptx.util import Inches

EMU_PER_INCH = 914400


def build(tmp_path, make_page, make_clip, *, full_pages: int = 3):
    plan = plan_for(full_pages=full_pages)
    images = {
        image_key("exc1", "partScore", 0): make_page("part.png", "Part"),
        **{
            image_key("exc1", "fullScore", page): make_page(f"f{page}.png", f"Full {page + 1}")
            for page in range(full_pages)
        },
    }
    clip = make_clip()
    deck = compose_deck(
        plan,
        images,
        {"exc1": Clip(path=clip, poster=make_page("poster.png", "Poster"))},
        tmp_path / "deck.pptx",
        title="NSO audition",
    )
    return plan, deck


@needs_ffmpeg
def test_deck_uses_the_slide_size_the_plan_was_laid_out_for(tmp_path, make_page, make_clip):
    plan, deck = build(tmp_path, make_page, make_clip)
    presentation = Presentation(str(deck))
    assert presentation.slide_width == Inches(SLIDE_W_IN)
    assert presentation.slide_height == Inches(SLIDE_H_IN)
    assert len(list(presentation.slides)) == len(plan)
    # the template ships type="screen4x3", which readers may honour over cx/cy
    assert presentation._element.sldSz.get("type") == "custom"


@needs_ffmpeg
def test_every_shape_stays_inside_the_slide(tmp_path, make_page, make_clip):
    _, deck = build(tmp_path, make_page, make_clip)
    presentation = Presentation(str(deck))
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, f"slide {index}: {shape.shape_type} off-slide"
            assert shape.left + shape.width <= Inches(SLIDE_W_IN) + 1
            assert shape.top + shape.height <= Inches(SLIDE_H_IN) + 1


@needs_ffmpeg
def test_the_clip_is_stored_once_for_the_whole_excerpt(tmp_path, make_page, make_clip):
    plan, deck = build(tmp_path, make_page, make_clip)
    with zipfile.ZipFile(deck) as archive:
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
    videos = [name for name in media if name.endswith(".mp4")]
    assert len(videos) == 1, f"clip duplicated per slide: {videos}"
    # part page + three full pages + the poster frame, each stored once
    assert len([name for name in media if name.endswith(".png")]) == 5
    assert len(plan) == 3


@needs_ffmpeg
def test_titles_and_pages_land_on_the_slides_they_belong_to(tmp_path, make_page, make_clip):
    _, deck = build(tmp_path, make_page, make_clip, full_pages=2)
    presentation = Presentation(str(deck))
    slides = list(presentation.slides)
    assert [shape.text_frame.text for shape in slides[0].shapes if shape.has_text_frame] == ["Mahler 1"]
    pictures = [shape for shape in slides[1].shapes if shape.shape_type == 13]
    assert len(pictures) == 2, "two portrait pages share the second slide"
