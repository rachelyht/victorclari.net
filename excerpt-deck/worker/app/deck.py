"""
Composes the .pptx from the SlidePlan the browser submits. The plan stores every box as a
fraction of the slide, so this is the same geometry the storyboard previews — it just has to be
written onto a slide of exactly SLIDE_W_IN x SLIDE_H_IN, or every shape lands off the slide.

python-pptx deduplicates media by hash, so the clip repeated on each slide of an excerpt is
stored once in the package.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
TITLE_BOX = {"x": 0.03, "y": 0.015, "w": 0.94, "h": 0.075}


@dataclass
class Clip:
    path: Path
    poster: Path | None = None
    start_sec: float | None = None
    end_sec: float | None = None
    trimmed: bool = True


def image_key(excerpt_id: str, source: str, page: int) -> str:
    return f"{excerpt_id}:{source}:{page}"


def _emu(box: dict) -> tuple[Emu, Emu, Emu, Emu]:
    return (
        Inches(box["x"] * SLIDE_W_IN),
        Inches(box["y"] * SLIDE_H_IN),
        Inches(box["w"] * SLIDE_W_IN),
        Inches(box["h"] * SLIDE_H_IN),
    )


def _time_label(seconds: float | None) -> str:
    if seconds is None:
        return ""
    total = max(0, round(seconds))
    return f"{total // 60}:{total % 60:02d}"


def _add_text(slide, box, text: str, *, size: int, bold: bool, color: str, align) -> None:
    left, top, width, height = box
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    frame.paragraphs[0].alignment = align


def compose_deck(
    plan: list[dict],
    images: dict[str, Path],
    clips: dict[str, Clip],
    output: Path,
    title: str = "Excerpt deck",
) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W_IN)
    presentation.slide_height = Inches(SLIDE_H_IN)
    # The default template's sldSz still claims screen4x3, which readers may prefer over cx/cy.
    presentation._element.sldSz.set("type", "custom")
    presentation.core_properties.title = title
    blank = presentation.slide_layouts[6]

    for entry in plan:
        slide = presentation.slides.add_slide(blank)

        if entry.get("showTitle", True) and entry.get("title"):
            _add_text(
                slide,
                _emu(TITLE_BOX),
                entry["title"],
                size=18,
                bold=True,
                color="222222",
                align=PP_ALIGN.LEFT,
            )

        for image in entry.get("images") or []:
            file = images.get(image_key(entry["excerptId"], image["source"], image["page"]))
            if not file:
                continue
            left, top, width, height = _emu(image)
            picture = slide.shapes.add_picture(str(file), left, top, width, height)
            picture.rotation = image.get("rotate") or 0

        box = entry.get("video")
        clip = clips.get(box["excerptId"]) if box else None
        if not clip:
            continue
        left, top, width, height = _emu(box)
        slide.shapes.add_movie(
            str(clip.path),
            left,
            top,
            width,
            height,
            poster_frame_image=str(clip.poster) if clip.poster else None,
            mime_type="video/mp4",
        )
        if not clip.trimmed and clip.start_sec is not None and clip.end_sec is not None:
            # Keynote's pptx importer ignores trim marks, so the window is written on the slide.
            label_top = min(top + height + Inches(0.05), Inches(SLIDE_H_IN - 0.35))
            _add_text(
                slide,
                (left, label_top, width, Inches(0.3)),
                f"{_time_label(clip.start_sec)} – {_time_label(clip.end_sec)}",
                size=11,
                bold=False,
                color="555555",
                align=PP_ALIGN.CENTER,
            )

    presentation.save(str(output))
    return output
